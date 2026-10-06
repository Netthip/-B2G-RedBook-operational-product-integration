# -*- coding: utf-8 -*-
"""BLOCKER 3 ของ #13 — ขอบเขตข้อมูลเข้า · ที่เขียนผล · ธงอ่านอย่างเดียว

🔴 `SYNTHETIC TEST FIXTURE` — ทุกสมุดงานในชุดนี้สร้างขึ้นในโฟลเดอร์ชั่วคราวระหว่างทดสอบ
ไม่มีชื่อหน่วยงาน รหัส หรือยอดเงินจริงใด ๆ · ไม่อ่านโฟลเดอร์ข้อมูลจริง
(Gift DECISION 6 ต.ค. 2569 · #13 ข้อ 2: พิสูจน์ด้วยข้อมูลสังเคราะห์ก่อน)
"""
import csv
import hashlib
import json
import os
import sys

import pytest
from openpyxl import Workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import build_registry as BR  # noqa: E402

NO_IDENTITY: dict = {}


def _wb(path, sheets=("README", "P1"), note="SYNTHETIC TEST FIXTURE"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheets[0]
    ws["A1"] = note
    for name in sheets[1:]:
        wb.create_sheet(name)["A1"] = 1.0
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


@pytest.fixture
def data_root(tmp_path):
    root = tmp_path / "data_root"
    _wb(str(root / "synthetic_case_a.xlsx"))
    _wb(str(root / "sub" / "mock_case_b.xlsx"), sheets=("P1",))
    _wb(str(root / "plain_case_c.xlsx"), sheets=("X", "Y"))
    return root


@pytest.fixture
def out_dir(tmp_path):
    return tmp_path / "out"


def test_fixture_dir_is_not_inside_a_git_work_tree(tmp_path):
    """ก่อนพิสูจน์อย่างอื่น — โฟลเดอร์ทดสอบเองต้องอยู่นอก git"""
    assert BR.find_git_work_tree(str(tmp_path)) is None


# ══════════════════════════════════════════ ข้อ 1 — ขอบเขตราก

def test_root_is_required_on_the_command_line(out_dir):
    with pytest.raises(SystemExit):
        BR.main(["--out", str(out_dir)])


def test_root_that_is_a_drive_root_is_refused(out_dir):
    drive = os.path.abspath(os.sep)
    with pytest.raises(BR.ScopeError):
        BR.check_root(drive)


def test_root_that_is_the_home_folder_is_refused():
    with pytest.raises(BR.ScopeError):
        BR.check_root(os.path.expanduser("~"))


def test_root_that_contains_this_repo_is_refused():
    parent_of_repo = os.path.dirname(BR._REPO_ROOT)
    with pytest.raises(BR.ScopeError):
        BR.check_root(parent_of_repo)


def test_discovery_never_leaves_the_root(tmp_path, data_root):
    """ไฟล์ข้างเคียง (นอกราก) ต้องไม่ถูกดูดเข้ามา"""
    _wb(str(tmp_path / "neighbour_project" / "synthetic_outside.xlsx"))
    found = BR.discover(str(data_root))
    assert len(found) == 3
    for p in found:
        assert BR._is_within(p, str(data_root))


def test_discovery_skips_hidden_folders_and_office_lock_files(data_root):
    _wb(str(data_root / ".hidden" / "synthetic_hidden.xlsx"))
    (data_root / "~$synthetic_case_a.xlsx").write_bytes(b"lock")
    names = {os.path.basename(p) for p in BR.discover(str(data_root))}
    assert "synthetic_hidden.xlsx" not in names
    assert "~$synthetic_case_a.xlsx" not in names


# ══════════════════════════════════════════ ข้อ 2 — ที่เขียนผล + provenance

def test_out_inside_a_git_work_tree_is_refused(tmp_path, data_root):
    fake_repo = tmp_path / "some_repo"
    (fake_repo / ".git").mkdir(parents=True)
    with pytest.raises(BR.ScopeError):
        BR.build(str(data_root), str(fake_repo / "registry"), identity=NO_IDENTITY)
    assert not (fake_repo / "registry").exists(), "ห้ามเขียนอะไรก่อนตรวจขอบเขต"


def test_out_inside_this_public_repo_is_refused(data_root):
    with pytest.raises(BR.ScopeError):
        BR.check_out(os.path.join(BR._REPO_ROOT, "registry"), str(data_root))


def test_out_under_the_root_is_refused(data_root):
    with pytest.raises(BR.ScopeError):
        BR.check_out(str(data_root / "out"), BR._real(str(data_root)))


def test_cli_refusal_exits_2_and_writes_nothing(tmp_path, data_root):
    fake_repo = tmp_path / "r"
    (fake_repo / ".git").mkdir(parents=True)
    rc = BR.main(["--root", str(data_root), "--out", str(fake_repo / "o")])
    assert rc == 2
    assert not (fake_repo / "o").exists()


def test_manifest_records_command_version_and_hash_of_every_output(data_root, out_dir):
    argv = ["--root", str(data_root), "--out", str(out_dir)]
    m = BR.build(str(data_root), str(out_dir), argv=argv, identity=NO_IDENTITY)
    on_disk = json.loads((out_dir / "build_manifest.json").read_text(encoding="utf-8"))
    assert on_disk == m
    assert m["builder_version"] == BR.BUILDER_VERSION
    assert m["argv"] == argv
    assert m["record_count"] == 3
    assert set(m["outputs"]) == {"registry.internal.json", "registry.internal.csv"}
    for name, meta in m["outputs"].items():
        data = (out_dir / name).read_bytes()
        assert hashlib.sha256(data).hexdigest() == meta["sha256"]
        assert len(data) == meta["bytes"]


def test_manifest_does_not_carry_the_real_root_path(data_root, out_dir):
    m = BR.build(str(data_root), str(out_dir), identity=NO_IDENTITY)
    blob = (out_dir / "build_manifest.json").read_text(encoding="utf-8")
    assert str(data_root) not in blob
    assert m["root_sha256"] == hashlib.sha256(
        BR._real(str(data_root)).encode("utf-8")).hexdigest()


def test_no_public_file_is_written_by_default(data_root, out_dir):
    BR.build(str(data_root), str(out_dir), identity=NO_IDENTITY)
    assert sorted(os.listdir(out_dir)) == [
        "build_manifest.json", "registry.internal.csv", "registry.internal.json"]


# ══════════════════════════════════════════ ข้อ 3 — ธงอ่านอย่างเดียวจากหลักฐาน

def test_read_only_flag_is_true_only_with_before_after_evidence(data_root, out_dir):
    BR.build(str(data_root), str(out_dir), identity=NO_IDENTITY)
    reg = json.loads((out_dir / "registry.internal.json").read_text(encoding="utf-8"))
    for r in reg["records"]:
        assert r["read_only_verified"] is True
        assert r["read_only_reason"] == "sha256+size+mtime_unchanged"


def test_unreadable_file_is_not_marked_read_only_verified(data_root, out_dir):
    (data_root / "synthetic_broken.xlsx").write_bytes(b"not a workbook")
    m = BR.build(str(data_root), str(out_dir), identity=NO_IDENTITY)
    reg = json.loads((out_dir / "registry.internal.json").read_text(encoding="utf-8"))
    broken = [r for r in reg["records"] if r["file_name"] == "synthetic_broken.xlsx"][0]
    assert broken["read_only_verified"] is False
    assert broken["read_only_reason"] == "read_failed"
    assert broken["record_id"] in m["read_only_failed"]


def test_a_file_that_changes_during_read_is_not_marked_verified(data_root, monkeypatch):
    path = str(data_root / "synthetic_case_a.xlsx")
    real_load = BR.load_workbook
    touched = {"done": False}

    def load_and_touch(p, *a, **k):
        wb = real_load(p, *a, **k)
        if not touched["done"]:
            with open(path, "ab") as f:
                f.write(b"\0")
            touched["done"] = True
        return wb

    monkeypatch.setattr(BR, "load_workbook", load_and_touch)
    rec = BR.build_record(path, BR._real(str(data_root)), 1)
    assert rec["read_only_verified"] is False
    assert rec["read_only_reason"] in ("file_changed_during_read", "read_failed")


def test_build_does_not_modify_any_source_file(data_root, out_dir):
    before = {p: hashlib.sha256(open(p, "rb").read()).hexdigest()
              for p in BR.discover(str(data_root))}
    BR.build(str(data_root), str(out_dir), identity=NO_IDENTITY)
    after = {p: hashlib.sha256(open(p, "rb").read()).hexdigest()
             for p in BR.discover(str(data_root))}
    assert before == after


# ══════════════════════════════════════════ BLOCKER 1 — ฉบับ preview กรองจริง

def test_public_preview_filters_rows_and_drops_fingerprints(data_root, out_dir):
    m = BR.build(str(data_root), str(out_dir), emit_public_preview=True,
                 identity=NO_IDENTITY)
    assert m["public_preview"] == "NOT CLEARED FOR PUBLICATION"
    text = (out_dir / "registry.public-preview.csv").read_text(encoding="utf-8-sig")
    first, rest = text.split("\n", 1)
    assert first == "# NOT CLEARED FOR PUBLICATION"
    rows = list(csv.DictReader(rest.splitlines()))
    reg = json.loads((out_dir / "registry.internal.json").read_text(encoding="utf-8"))
    publishable = {r["record_id"] for r in reg["records"] if r["publishable_to_github"]}
    not_publishable = {r["record_id"] for r in reg["records"]} - publishable
    assert not_publishable, "ชุดทดสอบต้องมีแถวที่ห้ามเผยแพร่อย่างน้อยหนึ่งแถว"
    assert {r["record_id"] for r in rows} == publishable
    for col in BR.FINGERPRINT_COLS:
        assert col not in rows[0]
    for r in reg["records"]:
        assert r["content_hash"] not in text


def test_the_withdrawn_public_registry_is_not_in_the_repo():
    """Gift DECISION 6 ต.ค. 2569 — ถอนออกจาก HEAD แบบ forward-only"""
    assert not os.path.exists(os.path.join(os.path.dirname(HERE), "registry.public.csv"))
